import pandas as pd

from PyQt6.QtWidgets import QWidget, QVBoxLayout, QLabel, QPushButton, QHBoxLayout, QTableWidget, QTableWidgetItem, QComboBox, QFileDialog, QDateEdit, QCheckBox
from PyQt6.QtCore import Qt, QDate
from PyQt6.QtGui import QFont
from app.database import get_db_session
from app.models import Buque, EtaCiudad, Tripulante, Vuelo, TripulanteVuelo, Transporte, TripulanteTransporte, Hotel, TripulanteHotel, Buque, Viaje
from app.controller.controllers import CITY_AIRPORT_CODES, CITY_TO_AIRPORT_CODES
from openpyxl.styles import PatternFill
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Side, Border
from datetime import datetime, time, timedelta
from sqlalchemy import func, and_, or_
from collections import defaultdict

class TransportesScreen(QWidget):
    def __init__(self, main_window):
        super().__init__()
        self.main_window = main_window
        self.setup_ui()

    def setup_ui(self):
        session = get_db_session()
        layout = QVBoxLayout(self)

        # Botón "Volver" para regresar a la pantalla anterior (Generación de Reportes)
        button_volver = QPushButton("Volver")
        button_volver.setStyleSheet("""
            font-size: 16px;  /* Tamaño de la letra */
            padding: 0px;    /* Elimina el espacio interno */
            line-height: 18px; /* Asegura que el texto no se corte verticalmente */
            text-align: center; /* Centra el texto */
        """)
        button_volver.setFixedWidth(80)
        button_volver.setFixedHeight(40)
        button_volver.clicked.connect(self.volver_a_opciones_programar)
        layout.addWidget(button_volver, alignment=Qt.AlignmentFlag.AlignLeft)

        self.label = QLabel("REQUERIMIENTO TRANSPORTES")
        self.label.setStyleSheet("""
            font-size: 40px;  /* Tamaño de fuente */
            font-weight: bold; /* Negrita */
            color: #00272d;    /* Color del texto */
            text-align: center; /* Centrar el texto horizontalmente */
            margin-bottom: 20px; /* Espacio debajo del título */
        """)

        layout.addWidget(self.label, alignment=Qt.AlignmentFlag.AlignCenter)  # Centrar el título
        layout.addSpacing(20)

        # Llenar el combo de ciudades desde la base de datos
        self.combo_ciudades = QComboBox()
        self.combo_ciudades.addItem("Ciudad")  # Agregar un valor por defecto

        ciudades = session.query(Transporte.city_in).distinct().all()  # Consulta para obtener las ciudades únicas
        for ciudad in ciudades:
            self.combo_ciudades.addItem(ciudad.city_in)
        layout.addWidget(self.combo_ciudades)

        self.check_fecha = QCheckBox("Habilitar filtro por fechas")
        self.check_fecha.setChecked(False)  # Inicialmente deshabilitado
        self.check_fecha.stateChanged.connect(self.actualizar_datos)  # Conectar evento de cambio de estado
        layout.addWidget(self.check_fecha)

        # Layout horizontal para las fechas
        layout_filtro = QHBoxLayout()

        # Selector de fecha de inicio
        self.date_start1 = QDateEdit()
        self.date_start1.setCalendarPopup(True)
        self.date_start1.setDate(QDate.currentDate())
        layout_filtro.addWidget(QLabel("Fecha inicio Pickup:"))
        layout_filtro.addWidget(self.date_start1)

        # Selector de fecha de fin
        self.date_end1 = QDateEdit()
        self.date_end1.setCalendarPopup(True)
        self.date_end1.setDate(QDate.currentDate())
        layout_filtro.addWidget(QLabel("Fecha fin Pickup:"))
        layout_filtro.addWidget(self.date_end1)

        layout.addLayout(layout_filtro)

        # Tabla para mostrar datos
        self.table_widget = QTableWidget()
        layout.addWidget(self.table_widget)

        # Botón para generar el Excel
        button_generar_excel = QPushButton("Generar Excel")
        button_generar_excel.setStyleSheet("""
            font-size: 18px;  /* Tamaño de la letra */
            padding: 0px;    /* Elimina el espacio interno */
            line-height: 18px; /* Asegura que el texto no se corte verticalmente */
            text-align: center; /* Centra el texto */
        """)
        button_generar_excel.setFixedHeight(45)
        button_generar_excel.setFixedWidth(245)
        button_generar_excel.clicked.connect(self.generar_excel_con_ciudad)
        layout.addWidget(button_generar_excel, alignment=Qt.AlignmentFlag.AlignRight)

        # Conectar el cambio en el QComboBox a un método
        self.combo_ciudades.currentTextChanged.connect(self.actualizar_datos)
        self.date_start1.dateChanged.connect(self.actualizar_datos)
        self.date_end1.dateChanged.connect(self.actualizar_datos)

        # Actualizar datos inicialmente
        self.actualizar_datos()

    def generar_excel_con_ciudad(self):
        ciudad_seleccionada = self.combo_ciudades.currentText()  # Obtener la ciudad seleccionada
        self.generar_excel(ciudad_seleccionada)  # Llamar a generar_excel con la ciudad seleccionada

    def actualizar_datos(self):
        ciudad_seleccionada = self.combo_ciudades.currentText()  # Obtiene la ciudad seleccionada
        if ciudad_seleccionada != "Ciudad":
            self.label.setText(f"REQUERIMIENTO TRANSPORTES EN {ciudad_seleccionada.upper()}")  # Actualiza el label
        else:
            self.label.setText(f"REQUERIMIENTO TRANSPORTES")  # Actualiza el label

        # Cargar datos en la tabla
        self.cargar_datos(ciudad_seleccionada)

    def toggle_fechas(self, state):
        enabled = state == Qt.CheckState.Checked  # Verificar si el checkbox está marcado
        self.date_start1.setEnabled(enabled)
        self.date_end1.setEnabled(enabled)
        # Deshabilitar el cuadro emergente si está desmarcado
        self.date_start1.setCalendarPopup(enabled)
        self.date_end1.setCalendarPopup(enabled)

    def cargar_datos(self, ciudad_seleccionada):
        session = get_db_session()
        ciudad_seleccionada = str(ciudad_seleccionada).lower()
        #print(ciudad_seleccionada)

        fecha_inicio = self.date_start1.date().toPyDate()
        fecha_fin = datetime.combine(self.date_end1.date().toPyDate(), time.max)

        # Construir la consulta de transporte
        transporte_necesario = (
            session.query(
                Tripulante.tripulante_id,
                Viaje.estado.label("Estado"),
                Tripulante.nombre.label("First_Name"),
                Tripulante.apellido.label("Last_Name"),
                Tripulante.nacionalidad.label("Nacionalidad"),
                Transporte.city_in.label("Ciudad_Transporte_in"),
                Transporte.place_in.label("Lugar_Transporte_in"),
                Transporte.city_end.label("Ciudad_Transporte_end"),
                Transporte.place_end.label("Lugar_Transporte_end"),
                TripulanteTransporte.date_pickup.label("Fecha_Pickup"),
                TripulanteTransporte.hours_pickup.label("Hora_Pickup")
            )
            .join(TripulanteTransporte, Tripulante.tripulante_id == TripulanteTransporte.tripulante_id)
            .join(Viaje, Tripulante.tripulante_id == Viaje.tripulante_id)
            .filter(Transporte.transporte_id == TripulanteTransporte.transporte_id,
                    Viaje.activo == 1)
            .filter(and_(func.lower(Transporte.city_in) == ciudad_seleccionada),
                    Transporte.transporte_id == TripulanteTransporte.transporte_id)
        )

        if self.check_fecha.isChecked():
            transporte_necesario = transporte_necesario.filter(
                TripulanteTransporte.date_pickup >= fecha_inicio,
                TripulanteTransporte.date_pickup <= fecha_fin
            )

        hotel_necesario = (
            session.query(
                Tripulante.tripulante_id,
                Tripulante.nombre.label("First_Name"),
                Tripulante.apellido.label("Last_Name"),
                Hotel.ciudad.label("Ciudad_Hotel"),
                Hotel.nombre.label("Nombre_Hotel"),
                TripulanteHotel.fecha_salida.label("Check_Out")
            )
            .join(TripulanteHotel, Tripulante.tripulante_id == TripulanteHotel.tripulante_id)
            .join(Viaje, Tripulante.tripulante_id == Viaje.tripulante_id)
            .filter(Hotel.hotel_id == TripulanteHotel.hotel_id)
            .filter(func.lower(Hotel.ciudad) == ciudad_seleccionada)
            .distinct()
        )

        if ciudad_seleccionada != "ciudad":    
            ciudad_seleccionada = CITY_AIRPORT_CODES.get(ciudad_seleccionada.upper())

        vuelo_necesario = (
            session.query(
                Tripulante.tripulante_id,
                Tripulante.nombre.label("First_Name"),
                Tripulante.apellido.label("Last_Name"),
                Vuelo.codigo.label("Codigo"),
                Vuelo.fecha.label("Fecha"),
                Vuelo.hora_salida.label("Hora_Salida"),
                Vuelo.hora_llegada.label("Hora_Llegada"),
                Vuelo.aeropuerto_salida.label("Aeropuerto_Salida"),
                Vuelo.aeropuerto_llegada.label("Aeropuerto_Llegada")
            )
            .join(TripulanteVuelo, Tripulante.tripulante_id == TripulanteVuelo.tripulante_id)
            .join(Viaje, Tripulante.tripulante_id == Viaje.tripulante_id)
            .filter(Vuelo.vuelo_id == TripulanteVuelo.vuelo_id)
            .filter(or_(
                func.lower(Vuelo.aeropuerto_salida) == ciudad_seleccionada.lower(),
                func.lower(Vuelo.aeropuerto_llegada) == ciudad_seleccionada.lower()
            ))
            .distinct()
        )

        buque_necesario = (
            session.query(
                Tripulante.tripulante_id,
                Buque.nombre.label("Nombre_Buque"),
                Buque.empresa.label("Owner"),  # Agregar la columna "empresa"
                EtaCiudad.eta.label("Eta")
            )
            .join(EtaCiudad, Tripulante.tripulante_id == EtaCiudad.tripulante_id)
            .join(Buque, EtaCiudad.buque_id == Buque.buque_id)
            .filter(Buque.buque_id == EtaCiudad.buque_id)
            .distinct()
        )

        transporte_necesario = transporte_necesario.all()
        hotel_necesario = hotel_necesario.all()
        vuelo_necesario = vuelo_necesario.all()
        buque_necesario = buque_necesario.all()

        # Organizar vuelos y transportes por tripulante_id
        transporte_dict = defaultdict(list)
        for transporte in transporte_necesario:
            transporte_dict[transporte.tripulante_id].append(transporte)

        vuelo_dict = defaultdict(list)
        for vuelo in vuelo_necesario:
            vuelo_dict[vuelo.tripulante_id].append(vuelo)

        hotel_dict = {hotel.tripulante_id: hotel for hotel in hotel_necesario}
        buque_dict = {buque.tripulante_id: (buque.Owner, buque.Nombre_Buque, buque.Eta) for buque in buque_necesario}

        headers = ["Estado", "Fecha Pickup", "Hora Pickup", "Nombre hotel", "Desde Ciudad", "Desde Lugar", "Hasta Ciudad", "Hasta Lugar", "Código Vuelo Llegada", "Date Llegada", "Hora Llegada", "Código Vuelo Salida", "Date Salida", "Fecha Salida", "Owner", "Nave", "ETA", "First Name", "Last Name", "Nacionalidad"]
        self.table_widget.setColumnCount(len(headers))
        self.table_widget.setHorizontalHeaderLabels(headers)
        self.table_widget.setRowCount(0)

        self.table_widget.resizeColumnsToContents()
        data_rows = []

        # Construir filas para cada tripulante con vuelos y transportes
        for tripulante_id, transportes in transporte_dict.items():
            vuelos = vuelo_dict.get(tripulante_id, [])
            hotel = hotel_dict.get(tripulante_id)
            owner, buque, eta = buque_dict.get(tripulante_id, ("", "", ""))

            for transporte in transportes:
                date_pickup = transporte.Fecha_Pickup
                tramo = f"{transporte.Lugar_Transporte_in}-{transporte.Lugar_Transporte_end}"
                city_select = CITY_AIRPORT_CODES.get(self.combo_ciudades.currentText(), "").lower()

                # Caso 'ATO-HOTEL'
                if 'ATO-HOTEL' == tramo:
                    vuelos_llegada = [v for v in vuelos if v.Aeropuerto_Llegada.lower() == city_select]
                    for vuelo in vuelos_llegada:
                        codigo = f"{str(vuelo.Codigo)} {CITY_TO_AIRPORT_CODES.get(vuelo.Aeropuerto_Salida)}-{CITY_TO_AIRPORT_CODES.get(vuelo.Aeropuerto_Llegada)}"
                        data_rows.append({
                            "estado": transporte.Estado,
                            "fecha_pickup": date_pickup,
                            "hora_pick_up": transporte.Hora_Pickup,
                            "nombre_hotel": hotel.Nombre_Hotel if hotel else "Sin hotel",
                            "ciudad_transporte_in": transporte.Ciudad_Transporte_in,
                            "lugar_transporte_in": transporte.Lugar_Transporte_in,
                            "ciudad_transporte_end": transporte.Ciudad_Transporte_end,
                            "lugar_transporte_end": transporte.Lugar_Transporte_end,
                            "codigo_vuelo": codigo,
                            "fecha_vuelo": vuelo.Fecha.date(),
                            "hora_salida": None,
                            "hora_llegada": vuelo.Hora_Llegada.time(),
                            "owner": owner,
                            "buque": buque,
                            "eta": eta,
                            "first_name": transporte.First_Name,
                            "last_name": transporte.Last_Name,
                            "nacionalidad": transporte.Nacionalidad
                        })

                # Caso 'HOTEL-ATO'
                elif 'HOTEL-ATO' == tramo:
                    vuelos_salida = [v for v in vuelos if v.Aeropuerto_Salida.lower() == city_select]
                    for vuelo in vuelos_salida:
                        codigo = f"{str(vuelo.Codigo)} {CITY_TO_AIRPORT_CODES.get(vuelo.Aeropuerto_Salida)}-{CITY_TO_AIRPORT_CODES.get(vuelo.Aeropuerto_Llegada)}"
                        tiempo_a_restar = timedelta(hours=2, minutes=30) if city_select == 'puq' else timedelta(hours=3, minutes=30)
                        hora_pick_up = (vuelo.Hora_Salida - tiempo_a_restar).time()
                        data_rows.append({
                            "estado": transporte.Estado,
                            "fecha_pickup": date_pickup,
                            "hora_pick_up": hora_pick_up,
                            "nombre_hotel": hotel.Nombre_Hotel if hotel else "Sin hotel",
                            "ciudad_transporte_in": transporte.Ciudad_Transporte_in,
                            "lugar_transporte_in": transporte.Lugar_Transporte_in,
                            "ciudad_transporte_end": transporte.Ciudad_Transporte_end,
                            "lugar_transporte_end": transporte.Lugar_Transporte_end,
                            "codigo_vuelo": codigo,
                            "fecha_vuelo": vuelo.Fecha.date(),
                            "hora_salida": vuelo.Hora_Salida.time(),
                            "hora_llegada": None,
                            "owner": owner,
                            "buque": buque,
                            "eta": eta,
                            "first_name": transporte.First_Name,
                            "last_name": transporte.Last_Name,
                            "nacionalidad": transporte.Nacionalidad
                        })

                elif 'ATO-NAVE' == tramo or 'NAVE-ATO' == tramo or 'NAVE-HOTEL' == tramo:
                    vuelos_llegada = [v for v in vuelos if v.Aeropuerto_Llegada.lower() == city_select]
                    for vuelo in vuelos_llegada:
                        codigo = f"{str(vuelo.Codigo)} {CITY_TO_AIRPORT_CODES.get(vuelo.Aeropuerto_Salida)}-{CITY_TO_AIRPORT_CODES.get(vuelo.Aeropuerto_Llegada)}"
                        data_rows.append({
                            "estado": transporte.Estado,
                            "fecha_pickup": date_pickup,
                            "hora_pick_up": transporte.Hora_Pickup,
                            "nombre_hotel": hotel.Nombre_Hotel if hotel else "Sin hotel",
                            "ciudad_transporte_in": transporte.Ciudad_Transporte_in,
                            "lugar_transporte_in": transporte.Lugar_Transporte_in,
                            "ciudad_transporte_end": transporte.Ciudad_Transporte_end,
                            "lugar_transporte_end": transporte.Lugar_Transporte_end,
                            "codigo_vuelo": codigo,
                            "fecha_vuelo": vuelo.Fecha.date(),
                            "hora_salida": None,
                            "hora_llegada": vuelo.Hora_Llegada.time(),
                            "owner": owner,
                            "buque": buque,
                            "eta": eta,
                            "first_name": transporte.First_Name,
                            "last_name": transporte.Last_Name,
                            "nacionalidad": transporte.Nacionalidad
                        })
        
        data_rows = [row for row in data_rows if row["fecha_pickup"] is not None]
                
        # Ordenar la lista de filas por `fecha_pickup`
        data_rows = sorted(data_rows, key=lambda x: x["fecha_pickup"])

        # Insertar filas ordenadas en la tabla
        self.table_widget.setRowCount(len(data_rows))
        for row, row_data in enumerate(data_rows):
            self.table_widget.setItem(row, 0, QTableWidgetItem(row_data["estado"]))
            self.table_widget.setItem(row, 1, QTableWidgetItem(str(row_data["fecha_pickup"])))
            self.table_widget.setItem(row, 2, QTableWidgetItem(str(row_data["hora_pick_up"])))
            self.table_widget.setItem(row, 3, QTableWidgetItem(row_data["nombre_hotel"]))
            self.table_widget.setItem(row, 4, QTableWidgetItem(row_data["ciudad_transporte_in"]))
            self.table_widget.setItem(row, 5, QTableWidgetItem(row_data["lugar_transporte_in"]))
            self.table_widget.setItem(row, 6, QTableWidgetItem(row_data["ciudad_transporte_end"]))
            self.table_widget.setItem(row, 7, QTableWidgetItem(row_data["lugar_transporte_end"]))
            if row_data["lugar_transporte_end"] != 'ATO':
                self.table_widget.setItem(row, 8, QTableWidgetItem(row_data["codigo_vuelo"]))
                self.table_widget.setItem(row, 9, QTableWidgetItem(str(row_data["fecha_vuelo"])))
                self.table_widget.setItem(row, 10, QTableWidgetItem(str(row_data["hora_llegada"])))
            else:
                self.table_widget.setItem(row, 11, QTableWidgetItem(row_data["codigo_vuelo"]))
                self.table_widget.setItem(row, 12, QTableWidgetItem(str(row_data["fecha_vuelo"])))
                self.table_widget.setItem(row, 13, QTableWidgetItem(str(row_data["hora_salida"])))
            self.table_widget.setItem(row, 14, QTableWidgetItem(row_data["owner"]))
            self.table_widget.setItem(row, 15, QTableWidgetItem(row_data["buque"]))
            self.table_widget.setItem(row, 16, QTableWidgetItem(str(row_data["eta"])))
            self.table_widget.setItem(row, 17, QTableWidgetItem(row_data["first_name"]))
            self.table_widget.setItem(row, 18, QTableWidgetItem(row_data["last_name"]))
            self.table_widget.setItem(row, 19, QTableWidgetItem(row_data["nacionalidad"]))
 

    def generar_excel(self, ciudad_seleccionada):
    # Crear un DataFrame con los datos de la tabla
        data = []
        for row in range(self.table_widget.rowCount()):
            row_data = []
            for column in range(self.table_widget.columnCount()):
                item = self.table_widget.item(row, column)
                row_data.append(item.text() if item else "")
            data.append(row_data)

        # Definir los nombres de las columnas
        headers = ["Estado", "Fecha Pickup", "Hora Pickup", "Nombre hotel", "Desde Ciudad", "Desde Lugar", "Hasta Ciudad", "Hasta Lugar", "Código Vuelo Llegada", "Date Llegada", "Hora Llegada", "Código Vuelo Salida", "Date Salida", "Fecha Salida", "Owner", "Nave", "ETA", "First Name", "Last Name", "Nacionalidad", "Transportista", "Anfitrión", "Vehículo", "Vehículo apoyo"]
        for row in data:
            while len(row) < len(headers):
                row.append("")  # Add empty strings to fill up to 24 columns
        df = pd.DataFrame(data, columns=headers)

        # Convertir "Hora Pick Up" a datetime para una ordenación correcta
        df['Hora Pickup'] = pd.to_datetime(df['Hora Pickup'], format='%H:%M:%S', errors='coerce')

        # Ordenar el DataFrame por "Hora Pick Up", "Lugar Pick Up" y "Nombre Hotel"
        df = df.sort_values(by=["Fecha Pickup", "Hora Pickup", "Desde Lugar", "Nombre hotel"])

        # Convertir "Hora Pick Up" de nuevo a solo hora para el Excel
        df['Hora Pickup'] = df['Hora Pickup'].dt.strftime('%H:%M:%S')

        file_name_parts = ["req_transporte"]
        if ciudad_seleccionada.lower() != "ciudad" and ciudad_seleccionada.lower() not in file_name_parts:
            file_name_parts.append(ciudad_seleccionada)
        if self.check_fecha.isChecked():
            file_name_parts.append(f"{self.date_start1.date().toPyDate()}_{self.date_end1.date().toPyDate()}")

        file_name = "_".join(file_name_parts) + ".xlsx"
        
        # Guardar en un archivo Excel
        file_path, _ = QFileDialog.getSaveFileName(
            self, 
            "Guardar archivo Excel", 
            file_name,  # Usar el nombre del archivo predefinido
            "Excel Files (*.xlsx);;All Files (*)"
        )
        
        if file_path:
            # Crear un nuevo libro de trabajo de Excel
            wb = Workbook()
            ws = wb.active

            # Escribir el texto final antes de la tabla
            ws.cell(row=1, column=1, value="Informe requerimientos transportes").font = Font(size=20, bold=True, underline="single")

            # Escribir "Date Pick Up" y "Ciudad" en negrita
            ws.cell(row=2, column=1, value="Date Pick Up").font = Font(bold=True)
            if self.check_fecha.isChecked():
                fecha_inicio = self.date_start1.date().toPyDate()
                fecha_fin = self.date_end1.date().toPyDate()
                ws.cell(row=2, column=2, value=f"({fecha_inicio}) - ({fecha_fin})")
            else:
                ws.cell(row=2, column=2, value="")

            ws.cell(row=3, column=1, value="Ciudad").font = Font(bold=True)
            ws.cell(row=3, column=2, value=ciudad_seleccionada)

            # Definir estilos de borde, relleno y alineación
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            header_fill = PatternFill(start_color='ADD8E6', end_color='ADD8E6', fill_type='solid')
            data_fill = PatternFill(start_color='FFFF99', end_color='FFFF99', fill_type='solid')

            # Escribir los encabezados de la tabla
            for col_num, col_name in enumerate(headers, 1):
                cell = ws.cell(row=5, column=col_num)
                cell.value = col_name
                cell.fill = header_fill
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

            # Escribir los datos y aplicar estilos
            for row_num, row_data in enumerate(df.values, start=6):
                for col_num, cell_value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_num, column=col_num)
                    cell.value = cell_value
                    cell.fill = data_fill
                    cell.border = thin_border

            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter  # Obtener la letra de la columna

                for cell in col[1:]:  # Empezar desde el índice 1, que corresponde a la fila 2 en Excel 
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass

                # Ajustar el ancho de la columna en base al contenido
                adjusted_width = max_length + 1
                ws.column_dimensions[column].width = adjusted_width


            # Guardar el archivo Excel con colores aplicados
            wb.save(file_path)

    def volver_a_opciones_programar(self):
        self.main_window.stacked_widget.setCurrentIndex(self.main_window.opciones_programar_index)