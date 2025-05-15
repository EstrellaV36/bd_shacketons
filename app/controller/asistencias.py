from datetime import datetime, timedelta
import pandas as pd
from sqlalchemy.orm import Session
from sqlalchemy import func
from app.models import Tripulante, TripulanteAsistencia
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

asistencia_columns = ['Proveedor SCL', 'Asistencia 1', 'Proveedor PUQ', 'Asistencia 2', 'Proveedor WPU', 'Asistencia 3']

class Asistencias:
    def __init__(self, db_session: Session):
        self.db_session = db_session

    def asistencias_main(self, file_path):
        try:
            # Leer datos de ambas hojas: ON y OFF
            excel_data_on = pd.read_excel(file_path, sheet_name='ON', header=None)
            excel_data_off = pd.read_excel(file_path, sheet_name='OFF', header=None)

            # Extraer asistencias de cada hoja
            asistencias_on = self._extract_assist(excel_data_on, start_row=2, column_range=slice(42, 48), column_names=asistencia_columns)
            asistencias_off = self._extract_assist(excel_data_off, start_row=2, column_range=slice(30, 36), column_names=asistencia_columns)

            asistencias_on.reset_index(drop=True, inplace=True)
            asistencias_off.reset_index(drop=True, inplace=True)

            # Normalizar datos de asistencia
            asistencias_on = self._normalize_assistance_dataframe(asistencias_on)
            asistencias_off = self._normalize_assistance_dataframe(asistencias_off)

            # print(f"ASISTENCIAS ON: {asistencias_on}")
            # print(f"ASISTENCIAS OFF: {asistencias_off}")

            return asistencias_on, asistencias_off
        except Exception as e:
            raise Exception(f"[Asistencias] Error al procesar el archivo: {e}")

    # def procesar_asistencias(self, file_path, tripulantes_on_df, asistencias_on_df, tripulantes_off_df, asistencias_off_df):
    #     errors = []
    #     errors_message = []
    #     try:
    #         print("Procesando asistencias para ON...")
    #         self._create_asistencias(file_path, tripulantes_on_df, asistencias_on_df, "ON", errors, errors_message)

    #         print("Procesando asistencias para OFF...")
    #         self._create_asistencias(file_path, tripulantes_off_df, asistencias_off_df, "OFF", errors, errors_message)

    #         print("Procesamiento de asistencias completado para ambas hojas.")
            
    #         return errors, errors_message
    #     except Exception as e:
    #         print(f"Error al procesar asistencias: {e}")

    def _create_asistencias(self, file_path, tripulantes_df, asistencias_df, state):
        errors = []
        errors_message = []

        try:
            if tripulantes_df.empty or asistencias_df.empty:
                print("No hay datos de tripulantes o asistencias para procesar.")
                return errors, errors_message

            # Abrir el archivo Excel una vez y crear un diccionario columna → letra
            workbook = load_workbook(file_path)
            sheet = workbook[state]
            column_letter_map = {
                col[0].value.strip(): get_column_letter(col[0].column)
                for col in sheet.iter_cols(1, sheet.max_column, 1, 1)
                if col[0].value
            }

            for (i, tripulante_row), (_, asistencia_row) in zip(tripulantes_df.iterrows(), asistencias_df.iterrows()):
                try:
                    if pd.isna(tripulante_row['Pasaporte']) or not tripulante_row['Pasaporte']:
                        print(f"Pasaporte vacío o nulo en fila {i}. Registro omitido.")
                        continue

                    tripulante = self.db_session.query(Tripulante).filter_by(pasaporte=tripulante_row['Pasaporte']).first()
                    if not tripulante:
                        print(f"No se encontró tripulante con pasaporte {tripulante_row['Pasaporte']} en la fila {i}.")
                        continue

                    asistencias_lista = [
                        str(asistencia_row.get('Asistencia 1', '')).strip().lower(),
                        str(asistencia_row.get('Asistencia 2', '')).strip().lower(),
                        str(asistencia_row.get('Asistencia 3', '')).strip().lower()
                    ]

                    proveedores_lista = [
                        asistencia_row.get('Proveedor SCL', None),
                        asistencia_row.get('Proveedor PUQ', None),
                        asistencia_row.get('Proveedor WPU', None)
                    ]

                    existing_asistencia = self.db_session.query(TripulanteAsistencia).filter_by(
                        tripulante_id=tripulante.tripulante_id
                    ).first()

                    asistencias_ciudades = [
                        ("asistencia scl", "Asistencia 1", "Proveedor SCL"),
                        ("asistencia puq", "Asistencia 2", "Proveedor PUQ"),
                        ("asistencia wpu", "Asistencia 3", "Proveedor WPU"),
                    ]

                    proveedor_bool = True

                    for idx, (asistencia_key, asistencia_column, proveedor_column) in enumerate(asistencias_ciudades):
                        proveedor_valor = str(proveedores_lista[idx]).strip().lower() if not pd.isna(proveedores_lista[idx]) else ""

                        if proveedor_valor == "no":
                            proveedor_bool = False
                            continue

                        if proveedor_valor in ["", "nan"]:
                            y = column_letter_map.get(proveedor_column, "?")
                            errors.append([i + 3, y])
                            errors_message.append(f"Proveedor {asistencia_key.split()[-1].upper()} faltante [{i + 3},{y}]")
                            proveedor_bool = False
                            continue

                        if asistencia_key not in asistencias_lista:
                            print(f"[ERROR ASISTENCIA] Asistencia {asistencia_key.split()[-1].upper()} faltante")
                            y = column_letter_map.get(asistencia_column, "?")
                            errors.append([i + 3, y])
                            errors_message.append(f"Asistencia inexistente en {asistencia_column} [{i + 3},{y}]")
                            proveedor_bool = False

                    if not proveedor_bool:
                        continue

                    if not existing_asistencia:
                        tripulante_asistencia = TripulanteAsistencia(
                            tripulante_id=tripulante.tripulante_id,
                            necesita_asistencia_scl='asistencia scl' in asistencias_lista,
                            necesita_asistencia_puq='asistencia puq' in asistencias_lista,
                            necesita_asistencia_wpu='asistencia wpu' in asistencias_lista,
                            proveedor_scl=proveedores_lista[0] if 'asistencia scl' in asistencias_lista else None,
                            proveedor_puq=proveedores_lista[1] if 'asistencia puq' in asistencias_lista else None,
                            proveedor_wpu=proveedores_lista[2] if 'asistencia wpu' in asistencias_lista else None
                        )
                        self.db_session.add(tripulante_asistencia)
                    else:
                        existing_asistencia.necesita_asistencia_scl = 'asistencia scl' in asistencias_lista
                        existing_asistencia.necesita_asistencia_puq = 'asistencia puq' in asistencias_lista
                        existing_asistencia.necesita_asistencia_wpu = 'asistencia wpu' in asistencias_lista
                        existing_asistencia.proveedor_scl = proveedores_lista[0] if 'asistencia scl' in asistencias_lista else None
                        existing_asistencia.proveedor_puq = proveedores_lista[1] if 'asistencia puq' in asistencias_lista else None
                        existing_asistencia.proveedor_wpu = proveedores_lista[2] if 'asistencia wpu' in asistencias_lista else None

                    self.db_session.commit()

                except Exception as row_error:
                    print(f"Error procesando asistencia en fila {i}: {row_error}")
                    self.db_session.rollback()

        except Exception as e:
            print(f"Error general al crear asistencias: {e}")
            self.db_session.rollback()

        return errors, errors_message

    def _extract_assist(self, data, start_row, column_range, column_names):
        try:
            data_block = []
            current_row = start_row

            while current_row < len(data):
                row_data = data.iloc[current_row, column_range]
                data_block.append(row_data)
                current_row += 1

            result_df = pd.DataFrame(data_block)

            if column_names:
                if len(column_names) != result_df.shape[1]:
                    raise ValueError(f"Length mismatch: Se esperaban {len(column_names)} columnas, pero se detectaron {result_df.shape[1]}")
                result_df.columns = column_names

            return result_df

        except Exception as e:
            raise Exception(f"Error al extraer asistencias: {e}")

    def _normalize_assistance_text(self, text):
        """Normaliza el texto de asistencia para manejar diferentes formatos."""
        if isinstance(text, str):
            words = sorted(text.lower().strip().split())
            return ' '.join(words)
        return text

    def _normalize_assistance_dataframe(self, df):
        """Normaliza las columnas relacionadas con las asistencias."""
        for col in ['Asistencia 1', 'Asistencia 2', 'Asistencia 3']:
            df[col] = df[col].apply(self._normalize_assistance_text)
        return df
    
def get_excel_column_letter(file_path, sheet_name, column_name):
        # Cargar el archivo y la hoja
        workbook = load_workbook(file_path)
        sheet = workbook[sheet_name]
        
        # Buscar la columna por nombre (suponiendo que los nombres están en la primera fila)
        for col in sheet.iter_cols(1, sheet.max_column, 1, 1):  # Iterar solo en la primera fila
            if col[0].value == column_name:
                # Devolver la letra de la columna
                return get_column_letter(col[0].column)
        
        raise ValueError(f"Columna con nombre '{column_name}' no encontrada en el archivo.")

def get_cell_value(file_path, sheet_name, row, column):
    # Cargar el archivo de Excel
    workbook = load_workbook(file_path, data_only=True)  # `data_only=True` para obtener el valor calculado en celdas con fórmulas
    sheet = workbook[sheet_name]

    # Obtener el valor de la celda
    cell_value = sheet.cell(row=row, column=column).value

    return cell_value